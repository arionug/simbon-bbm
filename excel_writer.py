import os
from datetime import datetime, date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "exports", "data-bon-bbm.xlsx")

HEADERS = [
    "Tanggal Nota", "Nomor Polisi", "Jenis Bahan Bakar", "Jumlah Liter",
    "Uang Sebanyak", "Terbilang", "Mengetahui", "Waktu Input",
]

NAMA_BULAN = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]


def _buat_sheet_bulan(wb, nama_sheet):
    """
    Membuat sheet baru untuk bulan tertentu
    lengkap dengan header dan styling.
    """
    ws = wb.create_sheet(title=nama_sheet)

    ws.append(HEADERS)

    header_fill = PatternFill(
        start_color="0F2D5A",
        end_color="0F2D5A",
        fill_type="solid"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    thin = Side(
        style="thin",
        color="CCCCCC"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(
            row=1,
            column=col_idx
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        cell.border = border

    # Lebar masing-masing kolom
    lebar_kolom = [
        14, 14, 18, 12,
        16, 32, 22, 20
    ]

    for i, lebar in enumerate(lebar_kolom, start=1):
        ws.column_dimensions[
            ws.cell(row=1, column=i).column_letter
        ].width = lebar

    ws.freeze_panes = "A2"

    return ws


def _buat_workbook_baru(nama_sheet):
    """
    Membuat workbook baru beserta sheet bulan pertama.
    """
    wb = Workbook()

    # Hapus sheet default
    ws_default = wb.active
    wb.remove(ws_default)

    ws = _buat_sheet_bulan(
        wb,
        nama_sheet
    )

    return wb, ws


def _ambil_tanggal(tanggal_nota):
    """
    Mengubah tanggal_nota menjadi object date/datetime
    jika masih berupa string.
    """

    if isinstance(tanggal_nota, (datetime, date)):
        return tanggal_nota

    if isinstance(tanggal_nota, str):

        format_tanggal = [
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y/%m/%d",
        ]

        for fmt in format_tanggal:
            try:
                return datetime.strptime(
                    tanggal_nota,
                    fmt
                )
            except ValueError:
                pass

    raise ValueError(
        "Format tanggal_nota tidak dikenali: "
        f"{tanggal_nota}"
    )


def _nama_sheet_bulan(tanggal_nota):
    """
    Menghasilkan nama sheet berdasarkan bulan dan tahun.

    Contoh:
    2026-09-08 -> September 2026
    """

    tanggal = _ambil_tanggal(tanggal_nota)

    return (
        f"{NAMA_BULAN[tanggal.month - 1]} "
        f"{tanggal.year}"
    )


def tambah_data_excel(row_data: dict):
    """
    Menambahkan data Nota Bon BBM ke Excel.

    Sheet otomatis dipisahkan berdasarkan bulan.

    row_data harus berisi:
        tanggal_nota
        nomor_polisi
        jenis_bbm
        jumlah_liter
        uang
        terbilang
        mengetahui
        waktu_input
    """

    os.makedirs(
        os.path.dirname(EXCEL_PATH),
        exist_ok=True
    )

    # Tentukan sheet berdasarkan tanggal nota
    nama_sheet = _nama_sheet_bulan(
        row_data["tanggal_nota"]
    )

    # ==========================================
    # Buka atau buat workbook
    # ==========================================

    if os.path.exists(EXCEL_PATH):

        wb = load_workbook(
            EXCEL_PATH
        )

        # Jika sheet bulan belum ada,
        # otomatis buat
        if nama_sheet in wb.sheetnames:

            ws = wb[nama_sheet]

        else:

            ws = _buat_sheet_bulan(
                wb,
                nama_sheet
            )

    else:

        wb, ws = _buat_workbook_baru(
            nama_sheet
        )

    # ==========================================
    # Data yang akan dimasukkan
    # ==========================================

    new_row = [
        row_data["tanggal_nota"],
        row_data["nomor_polisi"],
        row_data["jenis_bbm"],
        row_data["jumlah_liter"],
        row_data["uang"],
        row_data["terbilang"],
        row_data["mengetahui"],
        row_data["waktu_input"],
    ]

    # ==========================================
    # Cari baris kosong pertama
    # ==========================================

    target_row = None

    for row in range(
        2,
        ws.max_row + 1
    ):

        # Cek apakah seluruh kolom kosong
        if all(
            ws.cell(
                row=row,
                column=col
            ).value is None
            for col in range(
                1,
                len(HEADERS) + 1
            )
        ):

            target_row = row
            break

    # Kalau tidak ada baris kosong,
    # gunakan baris baru paling bawah
    if target_row is None:

        target_row = ws.max_row + 1

    # ==========================================
    # Styling data
    # ==========================================

    thin = Side(
        style="thin",
        color="E0E0E0"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    # ==========================================
    # Tulis data
    # ==========================================

    for col_idx, value in enumerate(
        new_row,
        start=1
    ):

        cell = ws.cell(
            row=target_row,
            column=col_idx
        )

        cell.value = value
        cell.border = border

        # Kolom Jumlah Liter dan Uang
        if col_idx in (4, 5):

            cell.alignment = Alignment(
                horizontal="right"
            )

    # ==========================================
    # Simpan
    # ==========================================

    wb.save(EXCEL_PATH)

    return EXCEL_PATH
